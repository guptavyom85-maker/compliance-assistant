from pathlib import Path
from datetime import date
from PIL import Image, ImageDraw, ImageFont

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
ASSET_DIR = ROOT / "tmp" / "docx_assets"
OUT_PATH = OUT_DIR / "Compliance_Assistant_Complete_Project_Guide.docx"
DIAGRAM_PATH = ASSET_DIR / "system_architecture.png"

NAVY = "17365D"
BLUE = "2F5597"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F3F5F7"
MID_GRAY = "D9D9D9"
DARK_GRAY = "4A4A4A"
WHITE = "FFFFFF"
BLACK = "000000"


def set_run_font(run, name="Arial", size=None, bold=None, color=BLACK, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=95, start=110, bottom=95, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=DARK_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_paragraph(doc, text="", style=None, before=0, after=5, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.keep_together = keep
    if text:
        r = p.add_run(text)
        set_run_font(r, size=10.5)
    return p


def add_lead_paragraph(doc, lead, text):
    p = add_paragraph(doc, after=6)
    r = p.add_run(lead)
    set_run_font(r, size=10.5, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = add_paragraph(doc, style=style, after=3)
    r = p.add_run(text)
    set_run_font(r, size=10.2)
    return p


def reset_numbering(doc):
    doc._guide_number = 0


def add_numbered(doc, text, level=0):
    if not hasattr(doc, "_guide_number"):
        reset_numbering(doc)
    doc._guide_number += 1
    p = add_paragraph(doc, after=3)
    p.paragraph_format.left_indent = Inches(0.28 + level * 0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    r = p.add_run(f"{doc._guide_number}.  ")
    set_run_font(r, size=10.2)
    r = p.add_run(text)
    set_run_font(r, size=10.2)
    return p


def add_code(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.right_indent = Inches(0.15)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), PALE_GRAY)
        p_pr.append(shd)
        r = p.add_run(line)
        set_run_font(r, name="Consolas", size=8.4, color="222222")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def add_table(doc, headers, rows, widths=None, font_size=8.7, header_fill=NAVY, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=font_size, bold=True, color=WHITE)
        if widths:
            cell.width = Inches(widths[idx])
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for cidx, value in enumerate(row_data):
            cell = row.cells[cidx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ridx % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            if alignments and cidx < len(alignments):
                p.alignment = alignments[cidx]
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
            if widths:
                cell.width = Inches(widths[cidx])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p


def build_architecture_diagram(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    w, h = 1800, 820
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    try:
        title_font = ImageFont.truetype("arialbd.ttf", 48)
        box_font = ImageFont.truetype("arialbd.ttf", 28)
        small_font = ImageFont.truetype("arial.ttf", 23)
    except OSError:
        title_font = ImageFont.load_default()
        box_font = ImageFont.load_default()
        small_font = ImageFont.load_default()

    draw.text((60, 30), "Compliance Assistant request and retrieval flow", fill=(0, 0, 0), font=title_font)

    boxes = [
        (70, 170, 350, 360, "Browser", "Django templates\nAuthenticated user"),
        (420, 170, 750, 360, "Django QA app", "Views, models, forms\nSQLite audit data"),
        (820, 170, 1120, 360, "Retriever", "Query embedding\nFAISS similarity"),
        (1190, 170, 1510, 360, "Grounded answer", "OpenRouter chat API\nPrompt with top chunks"),
        (1190, 500, 1510, 690, "Verification", "Paragraph ID check\nQueryLog record"),
        (470, 500, 800, 690, "Document pipeline", "PDF extraction\nClause chunking\nEmbedding and index"),
        (70, 500, 370, 690, "Regulatory PDFs", "Uploaded media\nDocument metadata"),
    ]

    for x1, y1, x2, y2, title, sub in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=(234, 242, 248), outline=(47, 85, 151), width=4)
        draw.text((x1 + 22, y1 + 22), title, fill=(0, 0, 0), font=box_font)
        y = y1 + 82
        for line in sub.split("\n"):
            draw.text((x1 + 22, y), line, fill=(45, 45, 45), font=small_font)
            y += 36

    def arrow(x1, y1, x2, y2):
        draw.line((x1, y1, x2, y2), fill=(23, 54, 93), width=8)
        import math
        angle = math.atan2(y2 - y1, x2 - x1)
        length = 24
        wing = 0.55
        p1 = (x2 - length * math.cos(angle - wing), y2 - length * math.sin(angle - wing))
        p2 = (x2 - length * math.cos(angle + wing), y2 - length * math.sin(angle + wing))
        draw.polygon([(x2, y2), p1, p2], fill=(23, 54, 93))

    arrow(350, 265, 420, 265)
    arrow(750, 265, 820, 265)
    arrow(1120, 265, 1190, 265)
    arrow(1350, 360, 1350, 500)
    arrow(1190, 595, 810, 595)
    arrow(370, 595, 470, 595)
    arrow(635, 500, 635, 365)
    arrow(420, 305, 355, 305)

    img.save(path, quality=95)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    if "Code" not in [s.name for s in styles]:
        code = styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code.font.size = Pt(8.4)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "Compliance Assistant Project Guide"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        set_run_font(r, size=8.3, color=DARK_GRAY)
    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_title_page(doc):
    doc.add_paragraph().paragraph_format.space_after = Pt(52)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(16)
    p_pr = p._p.get_or_add_pPr()
    p_border = p_pr.find(qn("w:pBdr"))
    if p_border is not None:
        p_pr.remove(p_border)
    r = p.add_run("RBI and SEBI Compliance Assistant")
    set_run_font(r, size=28, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("Complete Project Guide and Technical Assessment")
    set_run_font(r, size=15, bold=False, color=DARK_GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Repository reviewed")
    set_run_font(r, size=9.5, bold=True, color=DARK_GRAY)
    r = p.add_run("  C:\\Users\\Vyomkesh\\Desktop\\compliance-assiistant")
    set_run_font(r, size=9.5, color=DARK_GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Git revision")
    set_run_font(r, size=9.5, bold=True, color=DARK_GRAY)
    r = p.add_run("  508cc5c on main")
    set_run_font(r, size=9.5, color=DARK_GRAY)

    p = doc.add_paragraph()
    r = p.add_run("Assessment date")
    set_run_font(r, size=9.5, bold=True, color=DARK_GRAY)
    r = p.add_run("  3 October 2026")
    set_run_font(r, size=9.5, color=DARK_GRAY)

    doc.add_paragraph().paragraph_format.space_after = Pt(44)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Purpose")
    set_run_font(r, size=11, bold=True)
    p = add_paragraph(
        doc,
        "This guide explains the complete repository, current runtime state, data model, indexing and question-answering flows, setup and operating procedures, known limitations, and a prioritized path from course prototype to dependable compliance product.",
        after=7,
    )
    p = add_paragraph(
        doc,
        "The main conclusion is that the project has a coherent prototype architecture and a usable interface, but retrieval integrity, evaluation validity, corpus governance, authorization, and production configuration must be corrected before the system can support real compliance decisions.",
        after=0,
    )
    doc.add_page_break()


def add_guide_structure(doc):
    heading(doc, "Guide Structure", 1)
    add_paragraph(doc, "Word heading styles are used throughout so the Navigation pane can be used as a live table of contents.")
    sections = [
        "Executive Summary",
        "Project Purpose and Current State",
        "Repository and Folder Guide",
        "System Architecture",
        "End to End Workflows",
        "Django Application Guide",
        "RAG Package Guide",
        "Data Storage and Current Corpus",
        "Setup and Daily Operation",
        "Evaluation Methodology",
        "Security Privacy and Governance",
        "Known Issues and Technical Debt",
        "Recommended Capabilities and Product Extensions",
        "Prioritized Roadmap",
        "Troubleshooting and Runbook",
        "Appendices",
    ]
    for item in sections:
        add_bullet(doc, item)
    doc.add_page_break()


def add_executive_summary(doc):
    heading(doc, "Executive Summary", 1)
    add_paragraph(
        doc,
        "The repository implements a Django web application that accepts regulatory PDFs, splits their text into clause-like chunks, embeds those chunks with all-MiniLM-L6-v2, stores the vectors in FAISS, retrieves similar passages for a question, and sends the retrieved text to an OpenRouter-hosted language model. The interface also exposes document management, an audit log, and a small evaluation dashboard.",
    )
    add_paragraph(
        doc,
        "The design is understandable and appropriate for an academic RAG prototype. The code separates the web layer from retrieval logic, preserves source metadata, uses authenticated access for most operational pages, and records questions and answers. The current repository passes Django's normal system check and its database schema matches the models.",
    )
    add_paragraph(
        doc,
        "The live state nevertheless contains several correctness and governance problems. The FAISS index has 397 vectors while the database has 234 active chunks, leaving 163 stale vectors. The loaded corpus consists of a SEBI consultation paper and a chapter from an RBI financial stability report, while most seeded evaluation questions concern RBI digital lending. The evaluation logic does not compare generated answers with expected answers, and all seeded expected paragraph lists are empty. Its displayed accuracy therefore does not establish grounded answer quality.",
    )

    heading(doc, "Current Assessment", 2)
    rows = [
        ("Application structure", "Good prototype foundation", "Clear Django and RAG separation"),
        ("Runtime state", "Runs with reconstructed Python path", "The checked-in local virtual environment executable is broken"),
        ("Retrieval index", "Inconsistent", "397 vectors versus 234 active database chunks"),
        ("Corpus governance", "Not aligned", "Loaded documents do not match the stated digital-lending MVP"),
        ("Evaluation", "Not decision-grade", "Expected answer text is unused and paragraph expectations are empty"),
        ("Citation assurance", "Partial", "Checks paragraph identifiers only, not pages or claim support"),
        ("Security", "Development only", "Six Django deployment warnings and broad authenticated-user privileges"),
        ("Automated tests", "Absent", "Django discovered zero tests"),
    ]
    add_table(doc, ["Area", "Assessment", "Reason"], rows, widths=[1.55, 1.55, 3.75], font_size=8.8, alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 3)

    heading(doc, "Priority Actions", 2)
    reset_numbering(doc)
    for text in [
        "Rebuild the FAISS index from the active Chunk rows and make re-indexing replace vectors instead of appending duplicates.",
        "Replace the current evaluation scoring with retrieval metrics and answer-quality checks that use populated expected paragraph identifiers and expected answers.",
        "Load an approved corpus that matches the intended use case and record document type, version, effective period, supersession status, and authority level.",
        "Restrict upload, indexing, deletion, audit-log access, and evaluation runs by role; harden Django settings before any shared deployment.",
        "Add automated tests for chunking, index consistency, refusal behavior, citation verification, authorization, and the main web workflows.",
    ]:
        add_numbered(doc, text)


def add_project_state(doc):
    heading(doc, "Project Purpose and Current State", 1)
    heading(doc, "Intended Purpose", 2)
    add_paragraph(
        doc,
        "The project brief describes a regulatory compliance assistant for RBI and SEBI material. It is intended to answer in plain language using only retrieved regulatory passages, cite paragraph and page references, refuse questions not covered by the loaded corpus, and provide an evaluation trail that demonstrates grounding rather than merely a chatbot interface.",
    )
    add_paragraph(
        doc,
        "The original MVP narrative focused on the RBI Digital Lending Directions. It explicitly deferred hybrid retrieval, reranking, amendment linking, and automated regulatory monitoring. That scope is sensible for a course project, but the current local corpus no longer matches it.",
    )

    heading(doc, "Technology Stack", 2)
    add_table(
        doc,
        ["Layer", "Technology", "Role"],
        [
            ("Web", "Django 4.2.30", "Routing, authentication, forms, views, templates, and admin"),
            ("Database", "SQLite", "Documents, chunks, audit logs, gold questions, and evaluation results"),
            ("PDF parsing", "PyMuPDF 1.28.2", "Text extraction by PDF page"),
            ("Embeddings", "Sentence Transformers 6.1.0", "all-MiniLM-L6-v2 normalized 384-dimensional vectors"),
            ("Vector search", "FAISS 1.15.1", "IndexFlatIP cosine-like similarity over normalized embeddings"),
            ("Generation", "OpenAI Python client 3.20.0", "Calls OpenRouter's OpenAI-compatible chat endpoint"),
            ("Interface", "Django templates and Bootstrap 5.3.2", "Server-rendered dark-theme interface"),
        ],
        widths=[1.2, 2.1, 3.55],
        font_size=8.8,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 3,
    )

    heading(doc, "Live Snapshot", 2)
    add_table(
        doc,
        ["Item", "Observed value"],
        [
            ("First-party files", "48, excluding .git, .venv, and Python cache files"),
            ("Git branch and revision", "main at 508cc5c; one untracked nested repository named comp"),
            ("Documents", "2 total and 2 marked indexed"),
            ("Active chunks", "234"),
            ("FAISS vectors", "397"),
            ("Queries logged", "15"),
            ("Gold questions", "10"),
            ("Evaluation runs", "2, containing 20 results"),
            ("Users", "3 accounts; usernames and credentials are intentionally not reproduced"),
            ("Automated tests", "0 discovered"),
        ],
        widths=[2.25, 4.6],
        font_size=9,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 2,
    )


def add_repository_guide(doc):
    heading(doc, "Repository and Folder Guide", 1)
    add_paragraph(doc, "The source tree is small enough to understand as four logical areas: Django configuration, the qa application, the plain-Python RAG package, and runtime data or presentation assets.")
    add_code(doc, [
        "compliance-assiistant/",
        "  compliance_assistant/    Django project settings and root routes",
        "  qa/                      Models, views, forms, admin, migration, and seed command",
        "  rag/                     Chunking, embedding, FAISS, retrieval, LLM, and citation logic",
        "  templates/               Shared layout, login, Q&A, documents, logs, and evaluation pages",
        "  static/                  Custom CSS placeholder",
        "  media/documents/         Uploaded regulatory PDFs",
        "  vectorstore/             Persisted FAISS index and chunk ID mapping",
        "  db.sqlite3               Local application database",
        "  .env                     Local secrets and runtime values",
        "  .venv/                   Local environment, currently non-portable and broken",
        "  comp/                    Separate nested Git repository with only .gitattributes",
    ])

    heading(doc, "Important Root Files", 2)
    add_table(
        doc,
        ["Path", "Purpose", "Assessment"],
        [
            ("manage.py", "Django command entry point", "Standard and correct"),
            ("requirements.txt", "Broad dependency ranges", "Reproducibility would improve with a lock file"),
            ("README.md", "Repository introduction", "Only contains the repository name"),
            ("RAG_Compliance_Assistant_Project (1).md", "Original scope and architecture brief", "Useful but no longer matches the loaded corpus"),
            (".env.example", "Configuration template", "Covers API model, secret key, and debug mode"),
            (".gitignore", "Excludes secrets, environment, database, media, vector data, and IDE files", "Appropriate for source control"),
        ],
        widths=[2.25, 2.4, 2.25],
        font_size=8.7,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 3,
    )

    heading(doc, "Generated and Local Artifacts", 2)
    add_paragraph(
        doc,
        "The database, uploaded PDFs, vector store, environment file, IDE settings, virtual environment, and Python bytecode are excluded from the root repository by .gitignore. They are still part of this local folder and materially affect behavior. A new clone will not reproduce the current data or index unless these artifacts are transferred or rebuilt.",
    )
    add_paragraph(
        doc,
        "The comp directory is a separate Git repository with a single .gitattributes file and no remote. It is untracked by the main repository. It has no role in the application and should either be removed, documented, or deliberately integrated so it does not confuse future contributors.",
    )


def add_architecture(doc):
    heading(doc, "System Architecture", 1)
    add_paragraph(doc, "The system follows a conventional retrieval-augmented generation flow. Uploaded files are processed asynchronously only in the sense that the user initiates a separate indexing request; the actual work runs synchronously inside the web request.")
    build_architecture_diagram(DIAGRAM_PATH)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    shape = p.add_run().add_picture(str(DIAGRAM_PATH), width=Inches(6.9))
    shape._inline.docPr.set(
        "descr",
        "Architecture flow showing browser requests through Django, FAISS retrieval, OpenRouter generation, citation verification, query logging, and the PDF indexing pipeline.",
    )
    shape._inline.docPr.set("title", "Compliance Assistant architecture")
    p = doc.add_paragraph("Figure 1  Current application architecture")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(9)
    for r in p.runs:
        set_run_font(r, size=8.5, italic=True, color=DARK_GRAY)

    heading(doc, "Major Components", 2)
    add_table(
        doc,
        ["Component", "Primary files", "Responsibility"],
        [
            ("Project configuration", "compliance_assistant/settings.py and urls.py", "Environment, database, media, authentication redirects, RAG settings, and root routes"),
            ("Domain models", "qa/models.py", "Documents, extracted chunks, logs, gold questions, runs, and results"),
            ("Web workflows", "qa/views.py and qa/urls.py", "Question handling, document operations, logs, and evaluation"),
            ("Retrieval pipeline", "rag/pipeline.py and retriever.py", "Coordinates search, database lookup, generation, and citation checking"),
            ("Vector persistence", "rag/vectorstore.py", "Loads, searches, updates, and saves FAISS plus ID mapping"),
            ("Model integration", "rag/embedder.py and llm.py", "Local embeddings and remote OpenRouter completion"),
            ("Presentation", "templates and Bootstrap CDN", "Dark-theme server-rendered user interface"),
        ],
        widths=[1.45, 2.35, 3.1],
        font_size=8.5,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 3,
    )


def add_workflows(doc):
    heading(doc, "End to End Workflows", 1)
    heading(doc, "Document Upload and Indexing", 2)
    reset_numbering(doc)
    steps = [
        "An authenticated user opens /documents/upload/ and submits a title, file, regulator, status, and optional effective date.",
        "Django saves a Document row and places the file under media/documents/.",
        "The user clicks Index, which posts to /documents/<id>/index/.",
        "chunk_pdf reads each PDF page and creates a new chunk when a line begins with a recognized numeric, alphabetic, or Roman-numeral clause marker.",
        "Existing database chunks for the document are deleted and replacement Chunk rows are inserted.",
        "index_document embeds every current chunk and appends the vectors and chunk IDs to the global FAISS index.",
        "The document is marked indexed and its total chunk count is updated.",
    ]
    for item in steps:
        add_numbered(doc, item)
    add_lead_paragraph(doc, "Important consequence. ", "The vector index append operation does not remove the document's previous vectors during re-indexing. The current local mismatch shows this has already happened.")

    heading(doc, "Question Answering", 2)
    reset_numbering(doc)
    for item in [
        "The user submits a question on the authenticated Ask page.",
        "The application refuses to proceed if no document is marked indexed or no OpenRouter API key is configured.",
        "The query is embedded with the same all-MiniLM-L6-v2 model used for documents.",
        "FAISS returns the top inner-product matches; results below the 0.3 threshold are removed.",
        "The pipeline fetches matching Chunk rows from SQLite. Missing rows are silently discarded.",
        "The retrieved document title, paragraph identifier, page number, status, and text are assembled into the model context.",
        "The OpenRouter model is asked to answer only from that context, cite paragraph and page details, refuse unsupported questions, and add a legal disclaimer.",
        "A regular-expression checker extracts cited paragraph identifiers and verifies that each identifier appeared among the retrieved chunks.",
        "The answer, retrieved chunk IDs, scores, citation result, refusal flag, user, and response time are written to QueryLog.",
    ]:
        add_numbered(doc, item)

    heading(doc, "Evaluation", 2)
    reset_numbering(doc)
    for item in [
        "An authenticated user starts a batch run from /eval/run/.",
        "Each GoldQuestion is sent through the same live RAG and OpenRouter path.",
        "Unanswerable questions are correct only when the answer equals the exact not-found sentence.",
        "Answerable questions are graded from expected paragraph IDs, not from expected answer text.",
        "Partial retrieval counts as correct in the run-level accuracy calculation.",
        "Errors are stored as wrong results, and aggregate totals are displayed in the dashboard.",
    ]:
        add_numbered(doc, item)

    heading(doc, "Document Deletion", 2)
    add_paragraph(doc, "Deletion removes the document's current chunk vectors by reconstructing a new FAISS index from all retained positions, then deletes the Document row and cascades its Chunk rows. The uploaded PDF file itself is not deleted by Django when the model row is deleted, so orphaned media can remain on disk.")


def add_django_guide(doc):
    heading(doc, "Django Application Guide", 1)
    heading(doc, "Settings and Root Routing", 2)
    add_paragraph(doc, "compliance_assistant/settings.py loads .env at import time, uses SQLite, serves media directly only when DEBUG is true, and defines the RAG directory, API key, model name, top-k value, and similarity threshold. The current timezone is UTC, while the local user environment is Asia/Calcutta.")
    add_table(
        doc,
        ["Setting", "Current value", "Meaning"],
        [
            ("RAG_TOP_K", "5", "Maximum FAISS hits considered before database lookup"),
            ("RAG_CONFIDENCE_THRESHOLD", "0.3", "Minimum normalized inner-product score"),
            ("VECTORSTORE_DIR", "<project>/vectorstore", "FAISS and mapping persistence"),
            ("OPENROUTER_MODEL", "Environment or configured free-model default", "Primary remote generation model"),
            ("ALLOWED_HOSTS", "*", "Accepts every host header; unsuitable for production"),
            ("TIME_ZONE", "UTC", "Database display and timestamps use UTC-aware behavior"),
        ],
        widths=[1.8, 2.25, 2.85],
        font_size=8.7,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 3,
    )

    heading(doc, "Data Models", 2)
    add_table(
        doc,
        ["Model", "Key fields", "Operational role"],
        [
            ("Document", "title, file, regulator, status, effective_date, is_indexed", "One uploaded source and its declared regulatory status"),
            ("Chunk", "document, index, paragraph_id, text, page_number, metadata", "Retrieval unit linked to one document"),
            ("QueryLog", "user, question, answer, chunk IDs, scores, flags, time", "Audit record of each answered request"),
            ("GoldQuestion", "question, expected_answer, expected paragraph IDs, category", "Manual evaluation fixture"),
            ("EvalRun", "totals, response time, notes", "Batch-level summary"),
            ("EvalResult", "gold question, system answer, retrieval status, quality", "Per-question evaluation result"),
        ],
        widths=[1.25, 3.0, 2.65],
        font_size=8.5,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 3,
    )

    heading(doc, "Routes and Access", 2)
    add_table(
        doc,
        ["Route", "Method", "Access", "Purpose"],
        [
            ("/", "GET", "Public", "Home, examples, and corpus statistics"),
            ("/ask/", "GET POST", "Login", "Question form and RAG response"),
            ("/documents/", "GET", "Login", "Document inventory"),
            ("/documents/upload/", "GET POST", "Login", "Upload source PDF"),
            ("/documents/<id>/index/", "POST", "Login", "Parse and index document"),
            ("/documents/<id>/delete/", "POST", "Login", "Remove document and current vectors"),
            ("/query-log/", "GET", "Login", "View up to 100 recent queries from all users"),
            ("/eval/", "GET", "Login", "Evaluation dashboard"),
            ("/eval/run/", "POST", "Login", "Execute live evaluation"),
            ("/admin/", "GET POST", "Staff", "Django administration"),
        ],
        widths=[2.2, 0.8, 0.9, 2.9],
        font_size=8.2,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_paragraph(doc, "Only the administrative site checks staff status. Any authenticated user can upload, index, delete, inspect the global query log, and run evaluations. That is acceptable for a private classroom demo but too broad for an operational compliance tool.")

    heading(doc, "Templates and Interface", 2)
    add_paragraph(doc, "base.html supplies the dark theme, Bootstrap and Bootstrap Icons from jsDelivr, top navigation, sidebar, alert messages, and reusable styles for document status, answer boxes, and citation badges. Feature templates inherit it for the home screen, login, questions, documents, evaluation, and query history.")
    add_paragraph(doc, "static/css/custom.css is currently a placeholder; almost all custom styling is embedded in base.html. Moving the styles into static assets would improve caching, content-security-policy compatibility, and maintainability.")


def add_rag_guide(doc):
    heading(doc, "RAG Package Guide", 1)
    add_table(
        doc,
        ["File", "Public behavior", "Key limitation"],
        [
            ("chunker.py", "Extracts PDF text and starts chunks at clause-like lines", "Page-spanning chunks and short-chunk merging can make page or paragraph metadata inaccurate"),
            ("embedder.py", "Caches all-MiniLM-L6-v2 and returns normalized vectors", "Model name and embedding dimension are not stored with the index"),
            ("vectorstore.py", "Persists IndexFlatIP and a pickle ID list", "No locking, validation, atomic save, or idempotent document replacement"),
            ("retriever.py", "Applies top-k search and a fixed score threshold", "No keyword search, metadata filters, reranking, or stale-ID compensation"),
            ("llm.py", "Builds grounded prompt and uses one primary plus two fallback model names", "Provider errors are returned as user-visible answer text"),
            ("citation_checker.py", "Confirms cited paragraph IDs were retrieved", "Does not validate page numbers, claim coverage, quotation support, or document identity"),
            ("pipeline.py", "Joins retrieval, database chunks, generation, and checking", "A stale top-k result is dropped rather than replaced with the next valid hit"),
        ],
        widths=[1.35, 2.8, 2.75],
        font_size=8.3,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 3,
    )

    heading(doc, "Chunking Behavior", 2)
    add_paragraph(doc, "The clause detector recognizes a leading number such as 3 or 3.1, a lowercase alphabetic marker such as (a), or a Roman numeral such as (iv). It joins all subsequent non-matching lines until the next match. Chunks shorter than 50 characters are merged into the preceding chunk.")
    add_paragraph(doc, "This works reasonably for clean narrative regulations but is fragile for tables, headers, footers, annexures, multi-column reports, scanned pages, and clauses whose number is separated from their text. The NBFC PDF produced chunks from 50 to 6,067 characters; the SEBI paper produced chunks from 51 to 3,078 characters. Such variation affects retrieval quality.")
    add_paragraph(doc, "The parser records the PDF page where a chunk starts, not every page it covers. One observed SEBI chunk starts on PDF page 15 but includes the footer for page 16. The NBFC file's first PDF page is printed as report page 60, so a citation of page 1 from the application does not match the page number a reader sees in the document.")

    heading(doc, "Vector Search", 2)
    add_paragraph(doc, "Because document and query embeddings are normalized, IndexFlatIP behaves like exact cosine similarity. This is simple and transparent for a small corpus. It does not scale as efficiently as an approximate index, but 234 active chunks are far below the point where exact search is a performance concern.")
    add_paragraph(doc, "The more immediate concern is integrity. A test query returned both active and stale IDs with identical scores, showing that obsolete vectors remain searchable. Since the pipeline asks FAISS for only five hits before removing missing database rows, stale entries can crowd out valid context.")

    heading(doc, "Generation and Refusal", 2)
    add_paragraph(doc, "The generation prompt is deliberately strict about using only retrieved context and producing citations. Temperature is zero. The same prompt also requires an exact not-found sentence and a disclaimer at the end of the answer. Those instructions conflict when the context is insufficient: adding the disclaimer breaks the exact-string refusal check. The interface already shows a standing disclaimer, so refusal should be represented structurally rather than inferred from exact text.")


def add_data_storage(doc):
    heading(doc, "Data Storage and Current Corpus", 1)
    heading(doc, "SQLite State", 2)
    add_table(
        doc,
        ["Record type", "Count", "Notes"],
        [
            ("Documents", "2", "Both marked indexed and in force"),
            ("Chunks", "234", "IDs 1 to 71 and 235 to 397; IDs 72 to 234 are absent"),
            ("Query logs", "15", "7 citation-verified, 6 flagged, and 2 marked not found"),
            ("Gold questions", "10", "Most concern digital lending"),
            ("Evaluation runs", "2", "Displayed accuracies 10 percent and 70 percent"),
            ("Evaluation results", "20", "8 correct, 6 false refusals, and 6 wrong across both runs"),
        ],
        widths=[1.7, 1.0, 4.2],
        font_size=8.8,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    heading(doc, "Loaded Documents", 2)
    add_table(
        doc,
        ["Database title", "File and observed content", "Pages and chunks", "Governance concern"],
        [
            ("SEBI derivatives", "SEBI consultation paper on strengthening the index derivatives framework", "18 PDF pages and 71 chunks", "A consultation paper and draft circular are marked in force"),
            ("Shadow banking", "Chapter on financial institutions from the RBI Financial Stability Report June 2026", "60 PDF pages and 163 chunks", "A financial-stability report chapter is treated as an operative regulation"),
        ],
        widths=[1.3, 2.8, 1.25, 1.55],
        font_size=8.2,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 4,
    )
    add_paragraph(doc, "Both PDFs have extractable text on every page. Visual inspection showed professional source layouts, including multi-column narrative, tables, charts, footnotes, and printed page numbering. Those layouts are harder than plain regulatory circulars for the current line-based chunker. Extracted text also contains replacement characters for apostrophes and other glyphs.")

    heading(doc, "FAISS and Mapping Files", 2)
    add_paragraph(doc, "vectorstore/faiss_index.bin contains a 384-dimensional IndexFlatIP with 397 vectors. id_mapping.pkl contains 397 unique integer IDs from 1 through 397. SQLite contains only 234 of those IDs. The missing block 72 through 234 exactly matches the 163-chunk NBFC document, which indicates that an earlier version of that document was indexed and then its replacement vectors were appended.")
    add_paragraph(doc, "The mapping uses Python pickle. It should be considered trusted local state because loading an attacker-controlled pickle can execute code. A safer persistent mapping is JSON, a NumPy array, or a database table accompanied by integrity metadata.")

    heading(doc, "Secrets and Configuration", 2)
    add_paragraph(doc, "The local .env file contains values for OpenAI compatibility, Django, and OpenRouter. This guide intentionally records only variable names and whether they are configured; it does not reproduce secret values. The file is ignored by Git, which is correct.")


def add_setup(doc):
    heading(doc, "Setup and Daily Operation", 1)
    heading(doc, "Current Environment Note", 2)
    add_paragraph(doc, "The .venv folder contains installed packages, but its python.exe points to a uv-managed interpreter path that no longer exists. Direct commands through .venv fail. The project was validated by using a working Python 3.12 interpreter with the existing site-packages directory. For normal use, recreate the virtual environment instead of relying on this repair technique.")

    heading(doc, "Fresh Windows Setup", 2)
    reset_numbering(doc)
    add_numbered(doc, "Open PowerShell in the repository root.")
    add_numbered(doc, "Create and activate a fresh Python 3.12 virtual environment.")
    add_code(doc, [
        "py -3.12 -m venv .venv",
        ".\\.venv\\Scripts\\Activate.ps1",
    ])
    add_numbered(doc, "Install the application dependencies.")
    add_code(doc, ["python -m pip install --upgrade pip", "python -m pip install -r requirements.txt"])
    add_numbered(doc, "Copy the environment template and replace placeholder values.")
    add_code(doc, ["Copy-Item .env.example .env"])
    add_numbered(doc, "Create the database schema and an administrative account.")
    add_code(doc, ["python manage.py migrate", "python manage.py createsuperuser"])
    add_numbered(doc, "Seed the evaluation questions if the supplied set is appropriate for the loaded corpus.")
    add_code(doc, ["python manage.py seed_gold_questions"])
    add_numbered(doc, "Start the development server.")
    add_code(doc, ["python manage.py runserver"])
    add_numbered(doc, "Open http://127.0.0.1:8000/ and sign in before uploading or asking questions.")

    heading(doc, "Environment Variables", 2)
    add_table(
        doc,
        ["Variable", "Required", "Use"],
        [
            ("OPENROUTER_API_KEY", "Yes for answers and evaluation", "Authenticates remote generation requests"),
            ("OPENROUTER_MODEL", "No", "Overrides the primary model name"),
            ("DJANGO_SECRET_KEY", "Yes outside disposable local use", "Signs sessions and security-sensitive values"),
            ("DEBUG", "No", "Controls development behavior and media serving"),
            ("OPENAI_API_KEY", "Legacy fallback", "Used only if OPENROUTER_API_KEY is absent"),
        ],
        widths=[2.0, 2.0, 2.9],
        font_size=8.8,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 3,
    )

    heading(doc, "Daily Use", 2)
    for text in [
        "Upload only text-readable, approved regulatory documents and enter a precise title, authority, status, and effective date.",
        "Index a document once, then verify its chunk count and test several known questions before broader use.",
        "Review every answer's source passages. A green citation badge means cited paragraph identifiers were retrieved; it does not prove every statement is supported.",
        "Use the query log for review, but treat it as potentially sensitive because it stores user questions and generated answers.",
        "Run evaluation only after the gold set is aligned to the active corpus and expected paragraph identifiers are populated.",
    ]:
        add_bullet(doc, text)

    heading(doc, "Safe Index Maintenance", 2)
    add_paragraph(doc, "The repository does not include a safe rebuild command. Do not delete or replace vector files manually in a shared environment. Add a management command that builds a new index from active Chunk rows into temporary files, validates counts and dimensions, then atomically swaps the files. The same command should store the embedding model name and a corpus fingerprint.")


def add_evaluation(doc):
    heading(doc, "Evaluation Methodology", 1)
    heading(doc, "What the Dashboard Currently Measures", 2)
    add_paragraph(doc, "For answerable questions, a result is correct when all expected paragraph IDs are retrieved, partial when at least one is retrieved, and wrong when none are retrieved. When the expected list is empty, any non-refusal is counted correct. The expected_answer field is never compared with the generated answer. Partial retrieval also contributes fully to the run accuracy.")
    add_paragraph(doc, "For unanswerable questions, only the exact not-found sentence counts as a correct refusal. Model errors, fallback messages, or a not-found sentence accompanied by the requested disclaimer can be classified as an answer rather than a refusal.")

    heading(doc, "Why the Current Scores Are Not Valid Accuracy", 2)
    for text in [
        "All ten seeded questions have empty expected paragraph ID lists.",
        "The expected answer text is stored but unused.",
        "The current corpus does not contain the RBI Digital Lending Directions that most answerable questions target.",
        "A provider error string can be counted correct for an answerable question with no expected IDs.",
        "Citation verification and answer factuality are not part of the evaluation score.",
    ]:
        add_bullet(doc, text)

    heading(doc, "Recommended Evaluation Design", 2)
    add_table(
        doc,
        ["Metric", "Definition", "Why it matters"],
        [
            ("Recall at k", "Whether every required source passage appears in the first k valid hits", "Separates retrieval failure from generation failure"),
            ("Mean reciprocal rank", "Rank of the first required passage", "Measures how quickly relevant evidence appears"),
            ("Answer correctness", "Human rubric or calibrated model judge against expected answer and evidence", "Tests substance rather than response existence"),
            ("Faithfulness", "Every answer claim is supported by cited source text", "Detects grounded-sounding hallucinations"),
            ("Citation precision", "Every cited paragraph and page is accurate and relevant", "Validates the audit trail"),
            ("Refusal precision and recall", "Correct refusal on unsupported questions without false refusals", "Protects against both guessing and over-refusal"),
            ("Latency and cost", "Median and high-percentile time plus provider cost", "Supports operational decisions"),
        ],
        widths=[1.4, 3.0, 2.5],
        font_size=8.3,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 3,
    )
    add_paragraph(doc, "Create corpus-specific test sets with document versions frozen. Populate source paragraph and page expectations, include near-miss and superseded-rule questions, and record the exact embedding model, index fingerprint, generation model, prompt version, top-k, and threshold for every run.")


def add_security(doc):
    heading(doc, "Security Privacy and Governance", 1)
    heading(doc, "Deployment Findings", 2)
    add_paragraph(doc, "Django's standard check reports no functional configuration errors. The deployment check reports six warnings: no HSTS, no forced HTTPS, a weak or development-style secret key, non-secure session cookies, non-secure CSRF cookies, and DEBUG enabled. ALLOWED_HOSTS also accepts all hosts.")
    add_table(
        doc,
        ["Risk", "Current behavior", "Recommended control"],
        [
            ("Authorization", "Any logged-in user can upload, delete, run evaluation, and read the global query log", "Define viewer, analyst, corpus manager, evaluator, and administrator roles"),
            ("Sensitive queries", "Questions and answers are retained without a stated retention policy", "Add retention, access logging, export controls, and optional redaction"),
            ("File upload", "HTML accept filter only; no server-side PDF signature, type, or size validation", "Validate MIME signature, extension, page count, size, malware status, and parsing limits"),
            ("Prompt injection", "Retrieved PDF text is placed directly into the LLM prompt", "Treat documents as data, add injection tests, and constrain structured output"),
            ("External transfer", "Question and retrieved regulatory passages are sent to OpenRouter", "Document provider terms, data classification, residency, and approved-model policy"),
            ("Pickle mapping", "Mapping file is deserialized with pickle.load", "Use a non-executable format and validate checksums"),
            ("Concurrent writes", "FAISS and mapping files have no lock or atomic transaction", "Use a database-backed vector store or locked atomic rebuilds"),
            ("Source authority", "Consultation material can be marked in force", "Use document-type and authority-level workflows with approval"),
        ],
        widths=[1.35, 2.75, 2.8],
        font_size=8.2,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 3,
    )

    heading(doc, "Responsible Use", 2)
    add_paragraph(doc, "The interface correctly states that answers are informational and do not replace professional judgment. For a production compliance system, that disclaimer must be reinforced by workflow: source approval, answer review for high-risk questions, versioned evidence, traceable overrides, and clear escalation when the corpus is incomplete or conflicting.")


def add_issues(doc):
    heading(doc, "Known Issues and Technical Debt", 1)
    add_paragraph(doc, "The table prioritizes issues by their effect on correctness, governance, and operability. Priority 0 issues can invalidate answers or evaluation. Priority 1 issues block safe shared use. Priority 2 issues improve maintainability and user experience.")
    rows = [
        ("P0", "Stale vectors after re-index", "397 vectors but 234 active chunks; obsolete IDs consume top-k positions", "Replace per-document vectors or rebuild atomically from database"),
        ("P0", "Invalid evaluation score", "Expected answers unused and all expected paragraph lists empty", "Populate evidence labels and score retrieval, faithfulness, correctness, and refusal separately"),
        ("P0", "Corpus and gold-set mismatch", "Digital-lending questions are evaluated against derivatives and financial-stability documents", "Freeze an approved corpus and build a matching gold set"),
        ("P0", "Regulatory status misclassification", "Consultation paper and report chapter are marked in force", "Model document type, authority, dates, amendments, and supersession"),
        ("P0", "Page citations can be wrong", "Chunks can span pages and printed page numbers differ from PDF indices", "Track page spans and printed page labels per text segment"),
        ("P1", "Broad authenticated privileges", "All users can mutate corpus and inspect global logs", "Add role-based permissions and object-level access"),
        ("P1", "Non-transactional indexing", "Database chunks can be replaced before embeddings and files succeed", "Stage work and commit database plus index state atomically"),
        ("P1", "Exact-string refusal", "Disclaimer instruction conflicts with refusal detection", "Return typed JSON with answerable, answer, citations, and reason fields"),
        ("P1", "Weak citation verification", "Only cited paragraph IDs are checked", "Validate document, paragraph, page, spans, and claim support"),
        ("P1", "Deployment settings", "Six security warnings, DEBUG true, all hosts allowed", "Split development and production settings and enforce HTTPS controls"),
        ("P1", "No automated tests", "Zero tests discovered", "Add unit, integration, authorization, regression, and corpus tests"),
        ("P1", "Unvalidated uploads", "No server-side PDF or resource-limit validation", "Validate and quarantine uploads before parsing"),
        ("P1", "No file concurrency control", "Index readers and writers can race", "Add locks or move to a transactional vector database"),
        ("P2", "Broken local virtual environment", "Existing python.exe references a missing interpreter", "Recreate environment and document supported Python version"),
        ("P2", "Minimal README", "README contains only the project name", "Publish setup, architecture, operations, and contribution guidance"),
        ("P2", "Embedded CSS and CDN dependency", "Styles live in base template and require external CDN", "Move assets local and add a content security policy"),
        ("P2", "Unused form class", "QuestionForm is defined but the view reads request.POST directly", "Use the form for validation or remove it"),
        ("P2", "Orphaned uploaded files", "Deleting Document does not remove its PDF", "Add reviewed storage cleanup with recovery policy"),
        ("P2", "Provider errors shown as answers", "Exception text can be displayed and logged", "Use stable user messages and structured operational logs"),
    ]
    add_table(doc, ["Priority", "Issue", "Evidence and effect", "Recommended fix"], rows, widths=[0.55, 1.6, 2.65, 2.1], font_size=7.8, alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    heading(doc, "Validation Performed", 2)
    add_table(
        doc,
        ["Check", "Result"],
        [
            ("Django system check", "Passed with no normal configuration issues"),
            ("Migration drift", "No model changes detected"),
            ("Route smoke test", "Home and login returned 200; protected pages redirected to login"),
            ("Django deployment check", "Six security warnings"),
            ("Automated test discovery", "Zero tests found"),
            ("PDF extraction", "Both PDFs had extractable text on every page"),
            ("PDF visual sample", "Representative pages showed multi-column content, tables, charts, and printed page labels"),
            ("Vector integrity", "Index and mapping both had 397 entries; database had 234 active chunks"),
            ("Offline embedder load", "Cached all-MiniLM-L6-v2 loaded successfully"),
            ("Live generation", "Not invoked to avoid consuming the configured external API key during a repository audit"),
        ],
        widths=[2.25, 4.65],
        font_size=8.8,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 2,
    )


def add_extensions(doc):
    heading(doc, "Recommended Capabilities and Product Extensions", 1)
    add_paragraph(doc, "The strongest expansion path is to make source authority and evidence quality the product's center. The current Django shell can support that evolution without a full rewrite once the correctness issues are fixed.")

    heading(doc, "High Value Enhancements", 2)
    add_table(
        doc,
        ["Capability", "What it adds", "Value", "Effort"],
        [
            ("Hybrid retrieval", "BM25 or full-text search combined with embeddings", "Improves exact clause, circular number, and legal-term retrieval", "Medium"),
            ("Cross-encoder reranking", "Reorders a larger candidate set before generation", "Raises evidence precision without changing the interface", "Medium"),
            ("Evidence spans", "Stores exact character spans and page bounding boxes", "Creates defensible citations and source highlighting", "Medium to high"),
            ("Regulation version graph", "Links amendments, repeals, effective dates, and replaced clauses", "Answers which rule applied on a given date", "High"),
            ("Regulatory monitoring", "Scheduled discovery and review of new RBI and SEBI publications", "Keeps the corpus current with human approval", "High"),
            ("Document comparison", "Clause-level redline between versions", "Explains what changed and which controls are affected", "High"),
            ("Compliance obligation register", "Converts approved clauses into owners, deadlines, evidence, and status", "Moves from Q&A to operational compliance", "High"),
            ("Human review workflow", "Draft, review, approve, and publish answer states", "Supports high-risk use and auditability", "Medium"),
            ("Feedback and correction", "User ratings, disputed citations, and corrected answers", "Builds a measurable improvement loop", "Medium"),
            ("Exports", "Download answer packets with question, answer, sources, and version metadata", "Supports audit and committee reporting", "Low to medium"),
            ("API and integrations", "Authenticated REST API, case management, email, or chat connectors", "Embeds the assistant into existing work", "Medium to high"),
            ("OCR and table extraction", "Handles scanned circulars and complex financial tables", "Broadens usable source material", "Medium to high"),
            ("Multilingual queries", "Hindi and regional-language question normalization", "Improves accessibility while retaining source-language evidence", "Medium"),
            ("Analytics dashboard", "Topics, unresolved questions, latency, citation issues, and corpus gaps", "Shows where policy content or training is missing", "Medium"),
        ],
        widths=[1.55, 2.65, 2.2, 0.8],
        font_size=7.9,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )

    heading(doc, "Useful Product Directions", 2)
    for lead, text in [
        ("Compliance research assistant. ", "Keep the current question-and-answer experience, add authoritative source governance, precise citations, filters, and exports."),
        ("Regulatory change intelligence. ", "Monitor RBI and SEBI publications, compare versions, route changes for review, and alert affected teams."),
        ("Obligation and control mapper. ", "Convert approved regulatory clauses into obligations linked to policies, controls, evidence, owners, and tests."),
        ("Audit evidence workspace. ", "Preserve questions, approved answers, source snapshots, reviewer decisions, and downloadable evidence packages."),
        ("Training and policy support. ", "Use the same grounded corpus to create quizzes, role-specific guidance, and policy explanations with verified sources."),
    ]:
        add_lead_paragraph(doc, lead, text)

    heading(doc, "Features to Avoid Until the Foundation Is Fixed", 2)
    for text in [
        "Autonomous compliance decisions or automatic legal conclusions.",
        "Automatic ingestion that marks new documents authoritative without review.",
        "Large-scale model or interface work while retrieval and evaluation remain unreliable.",
        "Claims of legal accuracy based on the current evaluation dashboard.",
        "Broad integrations that expose query logs or regulatory context before access controls and retention rules exist.",
    ]:
        add_bullet(doc, text)


def add_roadmap(doc):
    heading(doc, "Prioritized Roadmap", 1)
    add_paragraph(doc, "The durations below are indicative for a small team and should be adjusted after confirming target users, deployment environment, and corpus size.")
    add_table(
        doc,
        ["Phase", "Indicative focus", "Exit criteria"],
        [
            ("Phase 0  Correctness", "Index rebuild, idempotent indexing, aligned corpus and gold set, structured refusals, reliable page spans", "Index count equals active chunk count; evaluation metrics are reproducible and meaningful"),
            ("Phase 1  Controlled pilot", "Role-based access, upload validation, tests, production settings, evidence spans, review workflow", "Named pilot users can safely operate a versioned corpus with auditable answers"),
            ("Phase 2  Retrieval quality", "Hybrid retrieval, reranking, metadata filters, threshold calibration, corpus diagnostics", "Target recall, faithfulness, citation precision, and refusal rates are met"),
            ("Phase 3  Regulatory operations", "Version graph, monitoring, change alerts, obligation mapping, exports", "Approved updates move from source discovery to owned compliance action"),
            ("Phase 4  Scale and integration", "Transactional vector service, background jobs, API, observability, high availability", "Service-level, privacy, recovery, and cost requirements are met"),
        ],
        widths=[1.25, 3.05, 2.6],
        font_size=8.4,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 3,
    )

    heading(doc, "First Ten Engineering Tasks", 2)
    reset_numbering(doc)
    for text in [
        "Add a rebuild_vector_index management command with atomic replacement and count validation.",
        "Change re-indexing to remove the document's existing vectors before adding replacements.",
        "Store embedding model, dimension, build time, corpus hash, and application revision beside the index.",
        "Create a corpus that matches the intended RBI digital-lending scope or revise the gold set and product framing to match the actual corpus.",
        "Populate expected paragraph IDs and add answer correctness plus faithfulness scoring.",
        "Return structured generation output and separate refusal state from display text.",
        "Track source page spans and printed page labels instead of one starting PDF page.",
        "Add Django permissions for corpus management, audit logs, and evaluation.",
        "Add tests for the RAG package and end-to-end web flows.",
        "Create production settings with secure cookies, HTTPS controls, restricted hosts, logging, and secret management.",
    ]:
        add_numbered(doc, text)


def add_runbook(doc):
    heading(doc, "Troubleshooting and Runbook", 1)
    add_table(
        doc,
        ["Symptom", "Likely cause", "Action"],
        [
            (".venv python cannot start", "Environment points to a missing uv interpreter", "Delete and recreate only the .venv directory, then reinstall requirements"),
            ("No documents have been indexed", "No Document is marked indexed", "Upload an approved PDF and complete indexing"),
            ("Model download or load fails", "Embedding model is absent or cache permissions fail", "Preload all-MiniLM-L6-v2 in the deployment environment and pin its revision"),
            ("OpenRouter key error", "OPENROUTER_API_KEY is missing", "Set the key in the protected runtime secret store or local .env"),
            ("Too few source passages", "Threshold is too high, stale IDs consumed top-k, or chunking is poor", "Check index integrity first, then evaluate retrieval and calibrate threshold"),
            ("Citation issue badge", "No recognized citation or a cited ID was not retrieved", "Review the answer and source passages; do not treat it as verified"),
            ("Wrong page citation", "Chunk spans pages or printed numbering differs", "Inspect the PDF and correct the page-span model before relying on citations"),
            ("Evaluation score is unexpectedly high", "Empty expected IDs count any non-refusal as correct", "Populate labels and use revised metrics"),
            ("Indexing partially fails", "Synchronous parse, database writes, embedding, or file save failed", "Keep the document unavailable, inspect logs, and rebuild from active chunks"),
            ("Text contains replacement glyphs", "PDF font encoding or extractor limitations", "Normalize text and compare extracted spans with rendered pages"),
        ],
        widths=[1.8, 2.45, 2.65],
        font_size=8.2,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] * 3,
    )

    heading(doc, "Release Checklist", 2)
    for text in [
        "Django system and deployment checks pass for the target environment.",
        "All migrations are applied and a rollback or backup exists.",
        "FAISS vector count equals the mapping count and every mapped chunk ID exists.",
        "The embedding model and index metadata match the application configuration.",
        "The approved corpus, document versions, and regulatory statuses are reviewed.",
        "Regression evaluation meets retrieval, faithfulness, citation, refusal, latency, and cost thresholds.",
        "Role permissions, audit-log visibility, retention, and provider data handling are approved.",
        "Upload limits, background-job behavior, logging, monitoring, backups, and recovery are tested.",
    ]:
        add_bullet(doc, text)


def add_appendices(doc):
    doc.add_page_break()
    heading(doc, "Appendix A Complete First Party File Catalog", 1)
    files = [
        (".env", "Local secret and runtime configuration; ignored by Git"),
        (".env.example", "Environment variable template"),
        (".gitignore", "Source-control exclusions"),
        (".vscode/settings.json", "Local editor Python settings"),
        ("comp/.gitattributes", "Only tracked file in the nested comp Git repository"),
        ("compliance_assistant/__init__.py", "Django project package marker"),
        ("compliance_assistant/asgi.py", "ASGI application entry point"),
        ("compliance_assistant/settings.py", "Project and RAG configuration"),
        ("compliance_assistant/urls.py", "Root routes, admin, auth, and development media"),
        ("compliance_assistant/wsgi.py", "WSGI application entry point"),
        ("db.sqlite3", "Local SQLite database"),
        ("manage.py", "Django command-line entry point"),
        ("media/documents/financialstatbility_report_for_NBFC.pdf", "RBI Financial Stability Report chapter used as a source"),
        ("media/documents/SEBI_Derivatives_In_depth.pdf", "SEBI index-derivatives consultation paper"),
        ("qa/__init__.py", "QA application package marker"),
        ("qa/admin.py", "Django admin registrations and displays"),
        ("qa/apps.py", "QA application configuration"),
        ("qa/forms.py", "Question and document upload forms"),
        ("qa/management/__init__.py", "Management package marker"),
        ("qa/management/commands/__init__.py", "Command package marker"),
        ("qa/management/commands/seed_gold_questions.py", "Seeds ten evaluation questions"),
        ("qa/migrations/__init__.py", "Migration package marker"),
        ("qa/migrations/0001_initial.py", "Initial schema migration"),
        ("qa/models.py", "Domain and evaluation models"),
        ("qa/urls.py", "Application URL patterns"),
        ("qa/views.py", "Web workflows"),
        ("RAG_Compliance_Assistant_Project (1).md", "Original project brief"),
        ("rag/__init__.py", "RAG public exports"),
        ("rag/chunker.py", "PDF clause extraction"),
        ("rag/citation_checker.py", "Paragraph citation validation"),
        ("rag/embedder.py", "Sentence Transformer embedding functions"),
        ("rag/llm.py", "OpenRouter generation and fallback models"),
        ("rag/pipeline.py", "Retrieval and indexing orchestration"),
        ("rag/retriever.py", "Thresholded top-k retrieval"),
        ("rag/vectorstore.py", "FAISS persistence and search"),
        ("README.md", "Minimal repository readme"),
        ("requirements.txt", "Python dependency ranges"),
        ("static/css/custom.css", "Placeholder custom stylesheet"),
        ("templates/base.html", "Shared dark-theme layout"),
        ("templates/qa/ask.html", "Question and answer interface"),
        ("templates/qa/document_list.html", "Document inventory and actions"),
        ("templates/qa/document_upload.html", "Upload interface"),
        ("templates/qa/eval_dashboard.html", "Evaluation history and details"),
        ("templates/qa/index.html", "Public home and statistics"),
        ("templates/qa/query_log.html", "Recent query table"),
        ("templates/registration/login.html", "Login interface"),
        ("vectorstore/faiss_index.bin", "Persisted 384-dimensional FAISS index"),
        ("vectorstore/id_mapping.pkl", "FAISS position to Chunk ID mapping"),
    ]
    add_table(doc, ["Path", "Role"], files, widths=[3.2, 3.7], font_size=8.0, alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    heading(doc, "Appendix B Glossary", 1)
    add_table(
        doc,
        ["Term", "Meaning in this project"],
        [
            ("Chunk", "A passage extracted from a PDF and used as the unit of retrieval"),
            ("Embedding", "A numeric vector representing the semantic content of text"),
            ("FAISS", "The local library and file format used for similarity search"),
            ("Gold question", "A manually prepared evaluation question with expected evidence and answer"),
            ("Grounding", "Restricting the answer to supplied source passages"),
            ("RAG", "Retrieval-Augmented Generation, which retrieves evidence before generating an answer"),
            ("Reranking", "A second model that reorders retrieved candidates by relevance"),
            ("Faithfulness", "The degree to which answer claims are supported by retrieved evidence"),
            ("Refusal", "A controlled not-found result when the corpus lacks sufficient support"),
            ("Vector drift", "Mismatch between persisted vectors and the current database or embedding configuration"),
        ],
        widths=[1.7, 5.2],
        font_size=8.8,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    heading(doc, "Appendix C Final Assessment", 1)
    add_paragraph(doc, "The repository is a credible academic prototype with a sensible separation of concerns and enough functionality to demonstrate document ingestion, semantic retrieval, grounded generation, citation display, and evaluation workflow. Its most important next step is not a larger model or more interface work. It is to make evidence state consistent and measurable.")
    add_paragraph(doc, "Once the index is rebuilt safely, the corpus and evaluation set are aligned, page-level evidence is reliable, and authorization is tightened, the project can become a useful internal research assistant. Version-aware regulation handling and change monitoring would then be the most valuable differentiators for a compliance audience.")


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    add_guide_structure(doc)
    add_executive_summary(doc)
    add_project_state(doc)
    add_repository_guide(doc)
    add_architecture(doc)
    add_workflows(doc)
    add_django_guide(doc)
    add_rag_guide(doc)
    add_data_storage(doc)
    add_setup(doc)
    add_evaluation(doc)
    add_security(doc)
    add_issues(doc)
    add_extensions(doc)
    add_roadmap(doc)
    add_runbook(doc)
    add_appendices(doc)
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_document()
